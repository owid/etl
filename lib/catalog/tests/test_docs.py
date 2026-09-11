"""Tests for the documentation rendered from metadata: codebook, sources, README, Excel workbook."""

import copy
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from owid.catalog import Dataset, DatasetMeta, License, Origin, Table
from owid.catalog.core import docs
from owid.catalog.core.datasets import EXCEL_MAX_ROWS

EI = Origin(
    producer="Energy Institute",
    title="Statistical Review of World Energy",
    date_published="2026-06-30",
    date_accessed="2026-07-02",
    url_main="https://www.energyinst.org/statistical-review/",
    citation_full="Energy Institute - Statistical Review of World Energy (2026).",
    license=License(name="© Energy Institute 2026", url="https://www.energyinst.org/terms"),
)
EI_COPY = copy.deepcopy(EI)
POPULATION = Origin(
    producer="Various sources",
    title="Population",
    attribution="Population based on various sources (2024)",
    date_published="2024-07-15",
    url_main="https://ourworldindata.org/population-sources",
    license=License(name="CC BY 4.0", url="https://creativecommons.org/licenses/by/4.0/"),
)


def make_table() -> Table:
    tb = Table(
        {
            "country": ["Spain", "Spain", "France"],
            "year": [2020, 2021, 2020],
            "hydro_energy_twh": [1.0, 2.0, 3.0],
            "hydro_energy_per_capita_kwh": [10.0, None, 30.0],
        },
        short_name="energy",
    )
    tb["hydro_energy_twh"].metadata.title = "Hydropower"
    tb["hydro_energy_twh"].metadata.description_short = "Energy from [hydropower](#dod:hydro)."
    tb["hydro_energy_twh"].metadata.unit = "terawatt-hours"
    tb["hydro_energy_twh"].metadata.short_unit = "TWh"
    tb["hydro_energy_twh"].metadata.origins = [EI]
    tb["hydro_energy_twh"].metadata.description_processing = "Converted from exajoules."
    tb["hydro_energy_per_capita_kwh"].metadata.title = "Hydropower per capita"
    tb["hydro_energy_per_capita_kwh"].metadata.unit = "kilowatt-hours per person"
    tb["hydro_energy_per_capita_kwh"].metadata.origins = [EI_COPY, POPULATION]
    tb["country"].metadata.title = "Country"
    tb["year"].metadata.title = "Year"
    return tb.set_index(["country", "year"])


def make_dataset(path: Path) -> Dataset:
    ds = Dataset.create_empty(
        path,
        DatasetMeta(
            namespace="energy",
            short_name="owid_energy",
            version="2026-09-10",
            title="Energy dataset",
            description="Key energy metrics.\n\n## Changelog\n\n- 2026-09-10: New methodology.",
            licenses=[License(name="CC BY 4.0", url="https://creativecommons.org/licenses/by/4.0/")],
        ),
    )
    ds.add(make_table())
    ds.save()
    return ds


def test_origin_label_defaults_to_producer_title_year():
    assert docs.origin_label(EI) == "Energy Institute – Statistical Review of World Energy (2026)"


def test_origin_label_prefers_attribution():
    assert docs.origin_label(POPULATION) == "Population based on various sources (2024)"


def test_origin_label_without_title_or_date():
    assert docs.origin_label(Origin(producer="Smil", title="Smil", date_published="latest")) == "Smil"


def test_codebook_uses_short_source_labels():
    codebook = make_table().codebook
    assert codebook.columns.tolist() == [
        "column",
        "title",
        "description",
        "unit",
        "source",
        "description_key",
        "description_processing",
    ]
    assert codebook["column"].tolist() == ["country", "year", "hydro_energy_twh", "hydro_energy_per_capita_kwh"]
    hydro = codebook.set_index("column").loc["hydro_energy_twh"]
    assert hydro["title"] == "Hydropower"
    assert hydro["description"] == "Energy from hydropower."
    assert hydro["unit"] == "terawatt-hours (TWh)"
    assert hydro["source"] == "Energy Institute – Statistical Review of World Energy (2026)"
    assert hydro["description_processing"] == "Converted from exajoules."
    assert hydro["description_key"] == ""
    per_capita = codebook.set_index("column").loc["hydro_energy_per_capita_kwh"]
    assert per_capita["source"] == (
        "Energy Institute – Statistical Review of World Energy (2026); Population based on various sources (2024)"
    )


def test_sources_deduplicates_and_orders_by_usage():
    sources = make_table().sources
    assert sources.columns.tolist() == docs.SOURCES_COLUMNS
    # The Energy Institute origin appears in two columns (as two equal objects); population in one.
    assert sources["label"].tolist() == [
        "Energy Institute – Statistical Review of World Energy (2026)",
        "Population based on various sources (2024)",
    ]
    ei = sources.set_index("label").loc["Energy Institute – Statistical Review of World Energy (2026)"]
    assert ei["url_main"] == "https://www.energyinst.org/statistical-review/"
    assert ei["license_name"] == "© Energy Institute 2026"
    assert ei["license_url"] == "https://www.energyinst.org/terms"
    assert ei["date_accessed"] == "2026-07-02"


def test_codebook_labels_match_sources_labels():
    tb = make_table()
    labels = {label for cell in tb.codebook["source"] for label in cell.split("; ") if label}
    assert labels == set(tb.sources["label"])


def test_readme_sections(tmp_path):
    ds = make_dataset(tmp_path / "owid_energy")
    readme = ds.readme(url="https://catalog.ourworldindata.org/energy/owid_energy/")
    assert readme.startswith("# Energy dataset\n")
    assert "https://catalog.ourworldindata.org/energy/owid_energy/" in readme
    assert "## Changelog" in readme
    assert "## How we process data at Our World in Data" in readme
    assert "### Hydropower\n\nEnergy from hydropower.\n" in readme
    assert "Column: `hydro_energy_twh`" in readme
    assert "Unit: terawatt-hours (TWh)" in readme
    assert "Column: `hydro_energy_twh`  \nUnit: terawatt-hours (TWh)  \nDate range: 2020–2021" in readme
    # The per-capita column has no 2021 value, so its range is shorter than the table's.
    assert "Column: `hydro_energy_per_capita_kwh`  \nUnit: kilowatt-hours per person  \nDate range: 2020–2020" in readme
    assert "Sources: Energy Institute – Statistical Review of World Energy (2026)" in readme
    assert "#### Notes on our processing step for this indicator\n\nConverted from exajoules." in readme
    assert "## Sources\n" in readme
    assert "### Energy Institute – Statistical Review of World Energy (2026)" in readme
    assert "Retrieved from: https://www.energyinst.org/statistical-review/" in readme
    assert "License: © Energy Institute 2026 (https://www.energyinst.org/terms)" in readme
    assert "Citation: Energy Institute - Statistical Review of World Energy (2026)." in readme
    assert "## License\n\nThis dataset is published under CC BY 4.0" in readme
    assert (
        "## How to cite this dataset\n\nEnergy Institute (2026); Population based on various sources (2024) – with "
        "minor processing by Our World in Data.\n\nFull citation: Energy Institute (2026); Population based on various "
        "sources (2024) – with minor processing by Our World in Data. “Energy dataset” [dataset]. Energy Institute, "
        "“Statistical Review of World Energy”; Various sources, “Population” [original data]. Retrieved from "
        "https://catalog.ourworldindata.org/energy/owid_energy/.\n"
    ) in readme
    assert "Last updated" not in readme
    # Nothing unrendered leaks through.
    assert "<%" not in readme and "<<" not in readme and "#dod:" not in readme


def test_readme_skips_jinja_templates(tmp_path):
    tb = make_table()
    tb["hydro_energy_twh"].metadata.title = "Hydropower << dimension >>"
    tb["hydro_energy_twh"].metadata.description_short = "<% if x %>Templated<% endif %>"
    ds = Dataset.create_empty(tmp_path / "ds", DatasetMeta(namespace="energy", short_name="ds"))
    ds.add(tb)
    ds.save()
    readme = ds.readme()
    assert "### hydro_energy_twh\n" in readme
    assert "Templated" not in readme
    assert "<%" not in readme and "<<" not in readme


def test_dataset_to_excel_sheets(tmp_path):
    ds = make_dataset(tmp_path / "owid_energy")
    path = tmp_path / "owid_energy.xlsx"
    ds.to_excel(path)
    workbook = openpyxl.load_workbook(path, read_only=True)
    assert workbook.sheetnames == ["data", "indicators", "sources", "readme"]
    data = pd.read_excel(path, sheet_name="data")
    assert data.columns.tolist() == ["country", "year", "hydro_energy_twh", "hydro_energy_per_capita_kwh"]
    assert len(data) == 3
    indicators = pd.read_excel(path, sheet_name="indicators")
    assert indicators["column"].tolist() == data.columns.tolist()
    sources = pd.read_excel(path, sheet_name="sources")
    assert len(sources) == 2
    readme = pd.read_excel(path, sheet_name="readme", header=None)
    assert readme.iloc[0, 0] == "# Energy dataset"


def test_dataset_to_excel_refuses_oversized_table(tmp_path, monkeypatch):
    ds = make_dataset(tmp_path / "owid_energy")
    monkeypatch.setattr("owid.catalog.core.datasets.EXCEL_MAX_ROWS", 2)
    assert EXCEL_MAX_ROWS > 2
    path = tmp_path / "owid_energy.xlsx"
    with pytest.raises(ValueError, match="row limit"):
        ds.to_excel(path)
    assert not path.exists()


def test_table_to_excel_includes_sources_sheet(tmp_path):
    path = tmp_path / "table.xlsx"
    make_table().to_excel(path)
    assert openpyxl.load_workbook(path, read_only=True).sheetnames == ["data", "metadata", "sources"]


def test_citations_follow_grapher_rules():
    origins = [
        EI,
        POPULATION,
        Origin(producer="Smil", title="Energy Transitions", date_published="2017", version_producer="2nd edition"),
        Origin(producer="Bolt and van Zanden", title="Maddison Project Database", date_published="2024-04-26"),
    ]
    assert docs.attribution_label(EI) == "Energy Institute (2026)"
    assert docs.attribution_label(POPULATION) == "Population based on various sources (2024)"
    # More than three attributions are shortened to the first one.
    assert (
        docs.citation_short(origins, processing_level="major")
        == "Energy Institute (2026) and other sources – with major processing by Our World in Data"
    )
    assert docs.citation_short(origins[:2], processing_level="minor") == (
        "Energy Institute (2026); Population based on various sources (2024) – with minor processing by Our World in Data"
    )
    long = docs.citation_long("Energy dataset", origins, processing_level="major", url="https://example.org/")
    assert long == (
        "Energy Institute (2026); Population based on various sources (2024); Smil (2017); Bolt and van Zanden (2024) – "
        "with major processing by Our World in Data. “Energy dataset” [dataset]. Energy Institute, “Statistical Review of "
        "World Energy”; Various sources, “Population”; Smil, “Energy Transitions 2nd edition”; Bolt and van Zanden, "
        "“Maddison Project Database” [original data]. Retrieved from https://example.org/."
    )
